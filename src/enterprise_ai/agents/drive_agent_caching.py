import logging
import os
from dataclasses import dataclass
from typing import Dict, List, Optional

from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
from strands import Agent

# Import VectorDB for caching
from enterprise_ai.storage.vector_store import MongoVectorStore

logger = logging.getLogger("drive_agent")


@dataclass
class OnboardingPlan:
    """Simple onboarding plan with ordered steps"""
    steps: List[str]
    estimated_duration: str
    is_accessible: bool = True
    is_mock: bool = False
    error_message: Optional[str] = None
    cached: bool = False  # Track if from cache
    source: str = "fresh"  # "cache", "drive", or "mock"


class DriveAgent:
    def __init__(self, credentials: Credentials = None, service_account_path: str = None):
        """Initialize DriveAgent with VectorDB caching"""
        self.logger = logging.getLogger("drive_agent")
        self.auth_method = "none"

        # Initialize Google Drive service
        self.drive_service = None
        self.is_authenticated = False

        if credentials:
            try:
                self.drive_service = build('drive', 'v3', credentials=credentials)
                self.is_authenticated = True
                self.auth_method = "user_oauth"
                self.logger.info("Using user OAuth credentials")
            except Exception as e:
                self.logger.error(f"Failed to initialize Drive: {e}")

        elif service_account_path and os.path.exists(service_account_path):
            try:
                from google.oauth2 import service_account
                sa_credentials = service_account.Credentials.from_service_account_file(
                    service_account_path,
                    scopes=['https://www.googleapis.com/auth/drive.readonly']
                )
                self.drive_service = build('drive', 'v3', credentials=sa_credentials)
                self.is_authenticated = True
                self.auth_method = "service_account"
                self.logger.info("Using service account credentials")
            except Exception as e:
                self.logger.error(f"Failed to load service account: {e}")

        self.agent = Agent(system_prompt=self._get_system_prompt())

        # NEW: Initialize VectorDB for caching
        try:
            self.vector_store = MongoVectorStore()
            self.logger.info("Vector store initialized for caching")
        except Exception as e:
            self.logger.warning(f"Vector store not available: {e}")
            self.vector_store = None

    def _get_system_prompt(self) -> str:
        return """You are an onboarding assistant. Given content from various documents:
        1. Analyze and understand the onboarding requirements
        2. Create a clear, ordered list of onboarding steps
        3. Ensure steps are concrete and actionable
        4. Include estimated time for each step
        Format your response as a numbered list of steps."""

    def create_onboarding_plan(self, drive_link: str) -> OnboardingPlan:
        """Create onboarding plan with VectorDB caching"""

        try:
            # Extract folder ID
            folder_id = self._extract_folder_id(drive_link)
            if not folder_id:
                raise ValueError("Could not extract folder ID from link")

            self.logger.info(f"📁 Processing folder: {folder_id}")

            # STEP 1: Check VectorDB cache first
            all_content = self._get_cached_content(folder_id)
            from_cache = len(all_content) > 0

            if from_cache:
                self.logger.info(f"Using {len(all_content)} cached documents from VectorDB")
                source = "cache"
            else:
                # STEP 2: Fetch from Drive and cache
                self.logger.info("No cache found, fetching from Google Drive...")
                all_content = self._fetch_and_cache_content(folder_id)
                source = "drive" if all_content else "fresh"

            if not all_content:
                return OnboardingPlan(
                    steps=[
                        "No content found",
                        "Possible reasons:",
                        "• Folder is empty",
                        "• No readable documents",
                        "• Permission denied"
                    ],
                    estimated_duration="N/A",
                    is_accessible=False,
                    cached=from_cache,
                    source=source
                )

            # STEP 3: Generate plan from content
            prompt = f"Create an onboarding plan based on these documents:\n\n{''.join(all_content)}"
            response = self.agent(prompt)

            # Parse response
            steps = [step.strip() for step in str(response).split('\n') if step.strip()]

            return OnboardingPlan(
                steps=steps if steps else ["No plan generated"],
                estimated_duration="1-2 weeks",
                is_accessible=True,
                is_mock=False,
                cached=from_cache,
                source=source
            )

        except Exception as e:
            self.logger.error(f"Error: {e}")
            return self._create_mock_onboarding_plan(drive_link)

    def _get_cached_content(self, folder_id: str) -> List[str]:
        """Retrieve cached documents from VectorDB"""
        if not self.vector_store or self.vector_store.collection is None:
            return []

        try:
            docs = self.vector_store.get_documents_by_folder(folder_id, limit=100)

            if docs:
                self.logger.info(f"Retrieved {len(docs)} cached documents")
                return [f"[{doc['source']}]\n{doc['content']}\n\n" for doc in docs]

            return []

        except Exception as e:
            self.logger.warning(f"Could not retrieve cache: {e}")
            return []

    def _fetch_and_cache_content(self, folder_id: str) -> List[str]:
        """Fetch from Google Drive and store in VectorDB"""
        if not self.is_authenticated or not self.drive_service:
            self.logger.warning("Not authenticated, cannot fetch from Drive")
            return []

        try:
            documents = self._get_folder_contents(folder_id)
            if not documents:
                self.logger.warning("Folder is empty")
                return []

            all_content = []
            stored_count = 0
            error_count = 0

            for doc in documents:
                try:
                    doc_name = doc.get('name', 'Unknown')
                    content = self._get_document_content(
                        doc['id'],
                        doc.get('mimeType'),
                        doc_name
                    )

                    if content and content.strip():
                        # Store in VectorDB
                        if self.vector_store:
                            doc_id = self.vector_store.store_document(
                                content=content,
                                source=doc_name,
                                folder_id=folder_id,
                                metadata={
                                    "mime_type": doc.get('mimeType'),
                                    "drive_id": doc['id']
                                }
                            )

                            if doc_id:
                                all_content.append(f"[{doc_name}]\n{content}\n\n")
                                stored_count += 1
                                self.logger.info(f"Cached: {doc_name}")
                            else:
                                self.logger.info(f"Could not cache: {doc_name}")
                        else:
                            all_content.append(f"[{doc_name}]\n{content}\n\n")
                    else:
                        self.logger.info(f"Skipped (empty): {doc_name}")

                except Exception as e:
                    error_count += 1
                    self.logger.warning(f"Error processing {doc.get('name')}: {e}")

            self.logger.info(f"📊 Stored {stored_count} docs, errors: {error_count}")
            return all_content

        except Exception as e:
            self.logger.error(f"Error fetching from Drive: {e}")
            return []

    def _extract_folder_id(self, drive_link: str) -> Optional[str]:
        """Extract folder ID from Google Drive link"""
        try:
            if '/folders/' in drive_link:
                return drive_link.split('/folders/')[-1].split('?')[0].split('#')[0]
            elif '/drive/folders/' in drive_link:
                return drive_link.split('/drive/folders/')[-1].split('?')[0].split('#')[0]
            return None
        except Exception as e:
            self.logger.error(f"Could not extract folder ID: {e}")
            return None

    def _get_folder_contents(self, folder_id: str) -> List[Dict]:
        """Get list of files in a Google Drive folder"""
        if not self.drive_service:
            return []

        try:
            results = self.drive_service.files().list(
                q=f"'{folder_id}' in parents and trashed=false",
                fields="files(id, name, mimeType)",
                pageSize=100
            ).execute()

            files = results.get('files', [])
            self.logger.info(f"Found {len(files)} files in folder")
            return files

        except Exception as e:
            self.logger.error(f"Error getting folder contents: {e}")
            return []

    def _get_document_content(self, file_id: str, mime_type: str = None, file_name: str = "Unknown") -> str:
        """Get content from a Google Drive document/file"""
        if not self.drive_service:
            return ""

        try:
            # Skip binary files
            if self._is_unsupported_binary_file(mime_type):
                self.logger.info(f"Skipping unsupported binary file: {file_name}")
                return ""

            # Handle Google Docs
            if mime_type == 'application/vnd.google-apps.document':
                request = self.drive_service.files().export_media(
                    fileId=file_id,
                    mimeType='text/plain'
                )
                content = request.execute()
                return self._safe_decode(content, file_name)

            # Handle Google Sheets
            elif mime_type == 'application/vnd.google-apps.spreadsheet':
                request = self.drive_service.files().export_media(
                    fileId=file_id,
                    mimeType='text/csv'
                )
                content = request.execute()
                return self._safe_decode(content, file_name)

            # Handle Word documents
            elif mime_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
                return self._extract_word_text(file_id, file_name)

            # Handle Excel files
            elif mime_type == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet':
                return self._extract_excel_text(file_id, file_name)

            # Handle text files
            elif self._is_text_file(mime_type):
                request = self.drive_service.files().get_media(fileId=file_id)
                content = request.execute()
                return self._safe_decode(content, file_name)

            else:
                self.logger.info(f"Skipping unsupported file: {file_name} ({mime_type})")
                return ""

        except HttpError as e:
            self.logger.warning(f"HTTP Error reading {file_name}: {e.resp.status}")
            return ""
        except Exception as e:
            self.logger.warning(f"Error reading {file_name}: {e}")
            return ""

    def _extract_word_text(self, file_id: str, file_name: str) -> str:
        """Extract text from Word document (.docx)"""
        try:
            import io

            from docx import Document

            request = self.drive_service.files().get_media(fileId=file_id)
            file_content = request.execute()

            doc = Document(io.BytesIO(file_content))
            text_content = []

            for para in doc.paragraphs:
                if para.text.strip():
                    text_content.append(para.text)

            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text for cell in row.cells]
                    text_content.append(" | ".join(row_text))

            result = "\n".join(text_content)
            self.logger.info(f"✅ Extracted {len(text_content)} items from {file_name}")
            return result if result.strip() else ""

        except ImportError:
            self.logger.warning("python-docx not installed")
            return ""

    def _extract_excel_text(self, file_id: str, file_name: str) -> str:
        """Extract text from Excel file (.xlsx)"""
        try:
            import io

            from openpyxl import load_workbook

            request = self.drive_service.files().get_media(fileId=file_id)
            file_content = request.execute()

            workbook = load_workbook(io.BytesIO(file_content))
            text_content = []

            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                text_content.append(f"Sheet: {sheet_name}")

                for row in sheet.iter_rows(values_only=True):
                    row_text = [str(cell) if cell is not None else "" for cell in row]
                    clean_row = [cell for cell in row_text if cell.strip()]
                    if clean_row:
                        text_content.append(" | ".join(clean_row))

            result = "\n".join(text_content)
            return result if result.strip() else ""

        except ImportError:
            self.logger.warning("openpyxl not installed")
            return ""

    def _is_unsupported_binary_file(self, mime_type: str) -> bool:
        """Check if file is binary that we can't process"""
        if not mime_type:
            return False

        unsupported_types = [
            'application/pdf',
            'image/',
            'video/',
            'audio/',
            'application/zip',
        ]

        return any(mime_type.startswith(ut) for ut in unsupported_types)

    def _is_text_file(self, mime_type: str) -> bool:
        """Check if file is text-based"""
        if not mime_type:
            return False

        text_types = [
            'text/',
            'application/json',
            'application/xml',
        ]

        return any(mime_type.startswith(tt) for tt in text_types)

    def _safe_decode(self, content: bytes, file_name: str = "Unknown") -> str:
        """Safely decode content with multiple encoding attempts"""
        if not isinstance(content, bytes):
            return str(content)

        encodings = ['utf-8', 'iso-8859-1', 'cp1252', 'utf-16', 'gbk', 'gb2312']

        for encoding in encodings:
            try:
                return content.decode(encoding)
            except (UnicodeDecodeError, AttributeError, LookupError):
                continue

        try:
            return content.decode('utf-8', errors='ignore')
        except Exception as e:
            self.logger.error(f"Could not decode {file_name}: {e}")
            return ""

    def _create_mock_onboarding_plan(self, drive_link: str) -> OnboardingPlan:
        """Create mock onboarding plan for demo"""
        mock_steps = [
            "📚 Week 1: Company Orientation",
            "  - Team introductions",
            "  - Review company policies",
            "  - Complete HR forms",
            "",
            "🛠️ Week 2: Technical Setup",
            "  - Install development tools",
            "  - Set up environment",
            "  - Access code repositories",
            "",
            "📖 Week 3: Learning",
            "  - Code review process",
            "  - Architecture overview",
            "  - First ticket assignment",
        ]

        return OnboardingPlan(
            steps=mock_steps,
            estimated_duration="3 weeks",
            is_accessible=False,
            is_mock=True,
            cached=False,
            source="mock"
        )

    def get_cache_stats(self) -> Dict:
        """Get statistics about cached documents"""
        if self.vector_store:
            return self.vector_store.get_stats()
        return {}

    def clear_cache(self, folder_id: str) -> int:
        """Clear cache for a specific folder"""
        if self.vector_store:
            deleted = self.vector_store.delete_folder_documents(folder_id)
            self.logger.info(f"Cleared {deleted} cached documents")
            return deleted
        return 0
